from __future__ import annotations

import hashlib
import io
import math
import re
import shutil
import time
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Literal, Protocol
from uuid import uuid4

import pymupdf
from docx import Document as WordDocument
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pydantic import BaseModel, Field

from workpilot.context import detect_prompt_injection
from workpilot.telemetry import TELEMETRY

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".md", ".txt"}
MAX_UNCOMPRESSED_OOXML_BYTES = 50 * 1024 * 1024


class KnowledgeError(ValueError):
    pass


class ObjectStore(Protocol):
    def put(self, key: str, content: bytes) -> None: ...

    def get(self, key: str) -> bytes: ...

    def delete(self, key: str) -> None: ...


class LocalObjectStore:
    """Filesystem object-store adapter for local/dev; keys never contain user filenames."""

    def __init__(self, root: Path) -> None:
        self.root = root.resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def put(self, key: str, content: bytes) -> None:
        path = self._path(key)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)

    def get(self, key: str) -> bytes:
        return self._path(key).read_bytes()

    def delete(self, key: str) -> None:
        path = self._path(key)
        path.unlink(missing_ok=True)
        parent = path.parent
        if parent != self.root and parent.exists() and not any(parent.iterdir()):
            parent.rmdir()

    def _path(self, key: str) -> Path:
        path = (self.root / key).resolve()
        if self.root not in path.parents:
            raise KnowledgeError("invalid object key")
        return path

    def clear(self) -> None:
        shutil.rmtree(self.root, ignore_errors=True)


class KnowledgeDocument(BaseModel):
    id: str
    tenant_id: str
    owner_id: str
    filename: str
    media_type: str
    size_bytes: int
    sha256: str
    visibility: Literal["private", "tenant"]
    status: Literal["PROCESSING", "READY", "FAILED"] = "PROCESSING"
    error: str | None = None
    chunk_count: int = 0
    object_key: str
    created_at: float = Field(default_factory=time.time)


class KnowledgeChunk(BaseModel):
    id: str
    document_id: str
    tenant_id: str
    owner_id: str
    visibility: Literal["private", "tenant"]
    ordinal: int
    text: str
    page: int | None = None
    sheet: str | None = None
    trust: Literal["untrusted"] = "untrusted"
    injection_markers: list[str] = Field(default_factory=list)
    embedding: list[float] = Field(exclude=True)


class Citation(BaseModel):
    document_id: str
    filename: str
    chunk_id: str
    page: int | None = None
    sheet: str | None = None
    preview: str
    score: float
    lexical_rank: int | None = None
    vector_rank: int | None = None
    trust: Literal["untrusted"] = "untrusted"
    injection_markers: list[str] = Field(default_factory=list)


@dataclass(frozen=True, slots=True)
class ParsedSection:
    text: str
    page: int | None = None
    sheet: str | None = None


class HashEmbeddingGateway:
    """Offline embedding baseline with the same provider boundary as a hosted model."""

    def __init__(self, dimensions: int = 128) -> None:
        self.dimensions = dimensions

    def embed(self, text: str) -> list[float]:
        vector = [0.0] * self.dimensions
        for token in _tokens(text):
            digest = hashlib.sha256(token.encode()).digest()
            index = int.from_bytes(digest[:4], "big") % self.dimensions
            vector[index] += 1.0 if digest[4] & 1 else -1.0
        norm = math.sqrt(sum(value * value for value in vector)) or 1.0
        return [value / norm for value in vector]


class KnowledgeService:
    def __init__(
        self,
        object_store: ObjectStore,
        *,
        max_file_bytes: int = 10 * 1024 * 1024,
        embedding_gateway: HashEmbeddingGateway | None = None,
    ) -> None:
        self.object_store = object_store
        self.max_file_bytes = max_file_bytes
        self.embedding_gateway = embedding_gateway or HashEmbeddingGateway()
        self.documents: dict[str, KnowledgeDocument] = {}
        self.chunks: dict[str, list[KnowledgeChunk]] = {}

    def ingest(
        self,
        *,
        tenant_id: str,
        owner_id: str,
        filename: str,
        media_type: str,
        content: bytes,
        visibility: Literal["private", "tenant"] = "private",
    ) -> KnowledgeDocument:
        extension = Path(filename).suffix.lower()
        self._validate_upload(extension, content)
        document_id = f"doc-{uuid4().hex[:16]}"
        object_key = f"{tenant_id}/{document_id}/original{extension}"
        document = KnowledgeDocument(
            id=document_id,
            tenant_id=tenant_id,
            owner_id=owner_id,
            filename=Path(filename).name,
            media_type=media_type or "application/octet-stream",
            size_bytes=len(content),
            sha256=hashlib.sha256(content).hexdigest(),
            visibility=visibility,
            object_key=object_key,
        )
        self.documents[document_id] = document
        self.object_store.put(object_key, content)
        started = time.perf_counter()
        try:
            sections = _parse(extension, content)
            chunks = self._chunk(document, sections)
            if not chunks:
                raise KnowledgeError("document contains no indexable text")
            self.chunks[document_id] = chunks
            document.status = "READY"
            document.chunk_count = len(chunks)
            TELEMETRY.metrics.counter("knowledge_documents_total", status="ready")
        except Exception as exc:
            document.status = "FAILED"
            document.error = str(exc)[:500]
            TELEMETRY.metrics.counter("knowledge_documents_total", status="failed")
        TELEMETRY.metrics.observe(
            "knowledge_index_duration_seconds",
            time.perf_counter() - started,
            extension=extension,
            status=document.status.lower(),
        )
        return document

    def list_documents(self, tenant_id: str, user_id: str) -> list[KnowledgeDocument]:
        return [
            item
            for item in self.documents.values()
            if _can_read(item.tenant_id, item.owner_id, item.visibility, tenant_id, user_id)
        ]

    def get(self, tenant_id: str, user_id: str, document_id: str) -> KnowledgeDocument:
        try:
            document = self.documents[document_id]
        except KeyError as exc:
            raise KeyError("document not found") from exc
        if not _can_read(
            document.tenant_id,
            document.owner_id,
            document.visibility,
            tenant_id,
            user_id,
        ):
            raise KeyError("document not found")
        return document

    def original(self, tenant_id: str, user_id: str, document_id: str) -> bytes:
        document = self.get(tenant_id, user_id, document_id)
        return self.object_store.get(document.object_key)

    def delete(self, tenant_id: str, user_id: str, document_id: str, *, is_admin: bool) -> None:
        document = self.get(tenant_id, user_id, document_id)
        if document.owner_id != user_id and not is_admin:
            raise PermissionError("only the owner or an admin can delete this document")
        self.object_store.delete(document.object_key)
        self.chunks.pop(document_id, None)
        self.documents.pop(document_id, None)

    def search(
        self,
        *,
        tenant_id: str,
        user_id: str,
        query: str,
        top_k: int = 5,
        mode: Literal["lexical", "vector", "hybrid_rrf"] = "hybrid_rrf",
    ) -> list[Citation]:
        candidates = [
            chunk
            for chunks in self.chunks.values()
            for chunk in chunks
            if _can_read(
                chunk.tenant_id,
                chunk.owner_id,
                chunk.visibility,
                tenant_id,
                user_id,
            )
        ]
        if not candidates:
            return []
        query_counts = Counter(_tokens(query))
        query_embedding = self.embedding_gateway.embed(query)
        lexical = sorted(
            candidates,
            key=lambda item: _lexical_score(query_counts, item.text),
            reverse=True,
        )
        vector = sorted(
            candidates,
            key=lambda item: _cosine(query_embedding, item.embedding),
            reverse=True,
        )
        lexical_rank = {item.id: rank for rank, item in enumerate(lexical, 1)}
        vector_rank = {item.id: rank for rank, item in enumerate(vector, 1)}
        if mode == "lexical":
            ranked = lexical
        elif mode == "vector":
            ranked = vector
        else:
            ranked = sorted(
                candidates,
                key=lambda item: 1 / (60 + lexical_rank[item.id]) + 1 / (60 + vector_rank[item.id]),
                reverse=True,
            )
        citations: list[Citation] = []
        for chunk in ranked[:top_k]:
            document = self.documents[chunk.document_id]
            score = 1 / (60 + lexical_rank[chunk.id]) + 1 / (60 + vector_rank[chunk.id])
            citations.append(
                Citation(
                    document_id=document.id,
                    filename=document.filename,
                    chunk_id=chunk.id,
                    page=chunk.page,
                    sheet=chunk.sheet,
                    preview=chunk.text[:300],
                    score=round(score, 6),
                    lexical_rank=lexical_rank[chunk.id],
                    vector_rank=vector_rank[chunk.id],
                    injection_markers=chunk.injection_markers,
                )
            )
        TELEMETRY.metrics.observe("retrieval_result_count", len(citations), mode=mode)
        return citations

    def _validate_upload(self, extension: str, content: bytes) -> None:
        if extension not in ALLOWED_EXTENSIONS:
            raise KnowledgeError("unsupported file type")
        if not content:
            raise KnowledgeError("empty files are not allowed")
        if len(content) > self.max_file_bytes:
            raise KnowledgeError(f"file exceeds {self.max_file_bytes} byte limit")
        if b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE" in content:
            raise KnowledgeError("malware signature detected")
        if extension == ".pdf" and not content.startswith(b"%PDF-"):
            raise KnowledgeError("file signature does not match PDF")
        if extension in {".docx", ".xlsx"}:
            if not content.startswith(b"PK"):
                raise KnowledgeError("file signature does not match OOXML")
            _validate_ooxml_archive(content)

    def _chunk(
        self, document: KnowledgeDocument, sections: list[ParsedSection]
    ) -> list[KnowledgeChunk]:
        output: list[KnowledgeChunk] = []
        ordinal = 0
        for section in sections:
            for text in _split_text(section.text):
                markers = detect_prompt_injection(text)
                output.append(
                    KnowledgeChunk(
                        id=f"{document.id}-chunk-{ordinal}",
                        document_id=document.id,
                        tenant_id=document.tenant_id,
                        owner_id=document.owner_id,
                        visibility=document.visibility,
                        ordinal=ordinal,
                        text=text,
                        page=section.page,
                        sheet=section.sheet,
                        injection_markers=markers,
                        embedding=self.embedding_gateway.embed(text),
                    )
                )
                ordinal += 1
        return output


def _parse(extension: str, content: bytes) -> list[ParsedSection]:
    if extension == ".pdf":
        pdf = pymupdf.open(stream=content, filetype="pdf")
        try:
            if pdf.page_count > 500:
                raise KnowledgeError("PDF exceeds 500 page limit")
            return [
                ParsedSection(page=index + 1, text=pdf[index].get_text("text"))
                for index in range(pdf.page_count)
            ]
        finally:
            pdf.close()
    if extension == ".docx":
        document = WordDocument(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )
        return [ParsedSection(text=text)]
    if extension == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sections: list[ParsedSection] = []
        try:
            for sheet in workbook.worksheets[:50]:
                rows: list[str] = []
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= 10_000:
                        raise KnowledgeError("worksheet exceeds 10000 row limit")
                    values = [str(value) for value in row if value is not None]
                    if values:
                        rows.append(" | ".join(values))
                sections.append(ParsedSection(text="\n".join(rows), sheet=sheet.title))
        finally:
            workbook.close()
        return sections
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise KnowledgeError("text files must be UTF-8") from exc
    if "\x00" in text:
        raise KnowledgeError("binary content is not allowed as text")
    return [ParsedSection(text=text)]


def _validate_ooxml_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1_000:
                raise KnowledgeError("OOXML archive contains too many entries")
            if sum(item.file_size for item in entries) > MAX_UNCOMPRESSED_OOXML_BYTES:
                raise KnowledgeError("OOXML archive expands beyond safety limit")
            if any(item.filename.lower().endswith("vbaproject.bin") for item in entries):
                raise KnowledgeError("macro-enabled documents are not allowed")
    except zipfile.BadZipFile as exc:
        raise KnowledgeError("invalid OOXML archive") from exc


def _split_text(text: str, size: int = 900, overlap: int = 120) -> list[str]:
    normalized = re.sub(r"[ \t]+", " ", text).strip()
    if not normalized:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(len(normalized), start + size)
        chunks.append(normalized[start:end])
        if end == len(normalized):
            break
        start = end - overlap
    return chunks


def _tokens(text: str) -> list[str]:
    return re.findall(r"[a-z0-9_]+|[\u4e00-\u9fff]", text.lower())


def _lexical_score(query: Counter[str], text: str) -> float:
    document = Counter(_tokens(text))
    return sum(min(count, document[token]) for token, count in query.items()) / (
        math.sqrt(sum(document.values())) or 1
    )


def _cosine(left: list[float], right: list[float]) -> float:
    return sum(a * b for a, b in zip(left, right, strict=True))


def _can_read(
    document_tenant: str,
    owner_id: str,
    visibility: str,
    tenant_id: str,
    user_id: str,
) -> bool:
    return document_tenant == tenant_id and (visibility == "tenant" or owner_id == user_id)

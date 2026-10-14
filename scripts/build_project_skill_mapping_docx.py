from __future__ import annotations

from pathlib import Path

import build_ai_agent_skills_docx as base
from build_newcomer_docx import add_inline, set_run_font
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "WorkPilot-专业技能项目映射与源码解析.md"
OUTPUT = ROOT / "docs" / "WorkPilot-专业技能项目映射与源码解析.docx"


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(58)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("WorkPilot 专业技能项目映射与源码解析")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    set_run_font(subtitle.add_run("根据简历五类技能逐项核验代码实现与能力边界"), 14.5)

    for label in (
        "分析对象  WorkPilot 0.2.0 当前仓库",
        "核验依据  源码  测试  数据库迁移  部署资产  评测报告",
        "交付目的  项目学习  面试讲解  简历校准",
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(7)
        set_run_font(paragraph.add_run(label), 10.5)

    intro = document.add_paragraph()
    intro.paragraph_format.left_indent = Inches(0.48)
    intro.paragraph_format.right_indent = Inches(0.48)
    intro.paragraph_format.space_before = Pt(24)
    intro.paragraph_format.line_spacing = 1.35
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        intro,
        "本报告逐项说明专业技能的含义、工程作用、WorkPilot 中的实现方式、代码证据和当前边界。"
        "对没有实际使用的技术明确标注，避免把概念了解写成项目经验。",
        size=11,
    )

    version = document.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version.paragraph_format.space_before = Pt(48)
    set_run_font(version.add_run("版本日期  2026 年 10 月 14 日"), 10)
    document.add_page_break()


def add_contents(document: Document, headings: list[str]) -> None:
    heading = document.add_heading("目录与使用方法", level=1)
    heading.paragraph_format.page_break_before = False
    paragraph = document.add_paragraph()
    add_inline(
        paragraph,
        "第 1 章给出项目定位和核验结论，第 2 至 6 章按简历分类逐项解释，第 7 章串联完整链路，"
        "第 8 至 10 章用于查证、校准简历和制定补强顺序，第 11 至 12 章用于面试复习。",
    )
    for item in headings:
        entry = document.add_paragraph()
        entry.paragraph_format.left_indent = Inches(0.18)
        entry.paragraph_format.first_line_indent = Inches(-0.18)
        entry.paragraph_format.space_after = Pt(4)
        add_inline(entry, item)
    document.add_page_break()


def add_header(section) -> None:
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run("WorkPilot 专业技能项目映射与源码解析"), 8.5)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D9D9D9")
    borders = OxmlElement("w:pBdr")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def build_document() -> None:
    base.SOURCE = SOURCE
    base.OUTPUT = OUTPUT
    base.add_cover = add_cover
    base.add_contents = add_contents
    base.add_header = add_header
    base.build_document()

    document = Document(OUTPUT)
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("代码证据"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.core_properties.title = "WorkPilot 专业技能项目映射与源码解析"
    document.core_properties.subject = "基于源码证据的五类专业技能逐项分析"
    document.core_properties.author = "WorkPilot"
    document.core_properties.keywords = (
        "WorkPilot, Python, FastAPI, Agent, RAG, MCP, Temporal, Evaluation"
    )
    document.core_properties.comments = "Generated from the maintained Markdown source."
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

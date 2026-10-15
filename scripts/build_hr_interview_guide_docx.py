from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_newcomer_docx as base_docx
from build_newcomer_docx import (
    add_inline,
    add_page_number,
    add_table,
    configure_styles,
    parse_table,
    set_run_font,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "WorkPilot-HR面试问题与回答手册.md"
OUTPUT = ROOT / "docs" / "WorkPilot-HR面试问题与回答手册.docx"
PAGE_BREAK_HEADINGS = {
    "15 面试前个人化工作表",
    "15.3 两个行为案例",
    "16 面试前检查清单",
    "18 证据来源",
}

base_docx.BODY_FONT = "Arial Unicode MS"
base_docx.EAST_ASIA_FONT = "Arial Unicode MS"
base_docx.MONO_FONT = "Arial Unicode MS"


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("WorkPilot HR 面试问题与回答手册"), 26, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    set_run_font(subtitle.add_run("互联网大厂 AI Agent 实习岗位准备材料"), 15)

    audience = document.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    audience.paragraph_format.space_after = Pt(8)
    set_run_font(audience.add_run("适用方向  Agent 应用开发  Agent 平台  企业 AI 工程"), 10.5)

    scope = document.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.paragraph_format.space_after = Pt(34)
    set_run_font(scope.add_run("内容范围  项目介绍  真实性核验  行为问题  岗位动机  回答边界"), 10.5)

    intro = document.add_paragraph()
    intro.paragraph_format.left_indent = Inches(0.52)
    intro.paragraph_format.right_indent = Inches(0.52)
    intro.paragraph_format.line_spacing = 1.35
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        intro,
        "本手册基于 WorkPilot 当前代码、测试和评测证据整理。示范回答用于帮助组织思路，"
        "不能替代候选人的真实经历。面试前必须把个人贡献、提交、测试、到岗时间和行为案例"
        "替换为本人能够核验的事实。",
        size=11,
    )

    version = document.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version.paragraph_format.space_before = Pt(50)
    set_run_font(version.add_run("版本日期  2026 年 10 月 15 日"), 10)
    document.add_page_break()


def add_contents(document: Document, headings: list[str]) -> None:
    heading = document.add_heading("目录与使用顺序", level=1)
    heading.paragraph_format.page_break_before = False
    paragraph = document.add_paragraph()
    add_inline(
        paragraph,
        "先阅读第 1 至 5 章统一项目口径，再使用第 6 至 13 章练习问题。"
        "面试前完成第 15 章个人化工作表和第 16 章检查清单。可在 Word 导航窗格按标题跳转。",
    )
    for item in headings:
        entry = document.add_paragraph()
        entry.paragraph_format.left_indent = Inches(0.2)
        entry.paragraph_format.first_line_indent = Inches(-0.2)
        entry.paragraph_format.space_after = Pt(3)
        add_inline(entry, item)
    document.add_page_break()


def add_header(section) -> None:
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run("WorkPilot HR 面试问题与回答手册"), 8.5)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D9D9D9")
    borders = OxmlElement("w:pBdr")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def keep_table_rows_together(table) -> None:
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))


def add_body(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    body_started = False
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()

        if not body_started:
            if lines[index].startswith("## "):
                body_started = True
            else:
                index += 1
                continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines: list[str] = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                table_lines.append(candidate)
                index += 1
            add_table(document, parse_table(table_lines))
            keep_table_rows_together(document.tables[-1])
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            markdown_level = len(heading_match.group(1))
            heading_text = heading_match.group(2)
            if markdown_level == 2:
                paragraph = document.add_heading(heading_text, level=1)
            else:
                style = "Manual Heading 2" if markdown_level == 3 else "Manual Heading 3"
                paragraph = document.add_paragraph(style=style)
                add_inline(paragraph, heading_text)
            if heading_text in PAGE_BREAK_HEADINGS:
                paragraph.paragraph_format.page_break_before = True
            index += 1
            continue

        ordered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if ordered:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.26)
            paragraph.paragraph_format.first_line_indent = Inches(-0.26)
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline(paragraph, f"{ordered.group(1)}.  {ordered.group(2)}")
            index += 1
            continue

        unordered = re.match(r"^-\s+(.+)$", stripped)
        if unordered:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, unordered.group(1))
            index += 1
            continue

        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip()[1:].strip())
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.34)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.keep_together = True
            add_inline(paragraph, " ".join(quote_lines))
            for run in paragraph.runs:
                run.italic = True
            continue

        if not stripped:
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", ">", "|", "- "))
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(paragraph, " ".join(paragraph_lines))


def build_document() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    headings = [line[3:].strip() for line in markdown.splitlines() if line.startswith("## ")]

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    # Keep worksheet writing lines comfortably clear of the page-number footer.
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    configure_styles(document)
    add_cover(document)
    add_contents(document, headings)
    add_body(document, markdown)

    for document_section in document.sections:
        add_header(document_section)
        add_page_number(document_section.footer.paragraphs[0])

    document.core_properties.title = "WorkPilot HR 面试问题与回答手册"
    document.core_properties.subject = "互联网大厂 AI Agent 实习岗位 HR 面试准备"
    document.core_properties.author = "WorkPilot"
    document.core_properties.keywords = "WorkPilot, HR, interview, AI Agent, internship"
    document.core_properties.comments = "Generated from a maintained Markdown source."
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

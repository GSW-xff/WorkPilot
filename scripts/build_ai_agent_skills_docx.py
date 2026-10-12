from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_newcomer_docx as base_docx
from build_newcomer_docx import (
    add_code_block,
    add_inline,
    add_page_number,
    add_table,
    configure_styles,
    parse_table,
    set_run_font,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "AI-Agent-专业技能系统讲解.md"
OUTPUT = ROOT / "docs" / "AI-Agent-专业技能系统讲解.docx"

# Arial Unicode MS is installed in the rendering environment and contains the
# Chinese glyphs required by this guide. The default font used by the older
# builder is not discovered by the bundled headless LibreOffice runtime.
base_docx.BODY_FONT = "Arial Unicode MS"
base_docx.EAST_ASIA_FONT = "Arial Unicode MS"
base_docx.MONO_FONT = "Arial Unicode MS"


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(68)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AI Agent 工程专业技能系统讲解")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    set_run_font(subtitle.add_run("从零散接触到系统掌握的学习手册"), 15)

    audience = document.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    audience.paragraph_format.space_after = Pt(8)
    set_run_font(audience.add_run("适用对象  AI Agent 与后端工程方向求职者"), 10.5)

    scope = document.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.paragraph_format.space_after = Pt(36)
    set_run_font(
        scope.add_run("覆盖范围  后端工程  Agent  RAG  评测稳定性  LLM 基础"),
        10.5,
    )

    intro = document.add_paragraph()
    intro.paragraph_format.left_indent = Inches(0.55)
    intro.paragraph_format.right_indent = Inches(0.55)
    intro.paragraph_format.line_spacing = 1.35
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        intro,
        "本手册逐项解释简历中的专业技能，说明工作机制、工程场景、常见误区和实践标准；"
        "随后按五类能力梳理完整技术链路，并给出学习顺序、项目练习、面试追问和简历措辞建议。",
        size=11,
    )

    version = document.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version.paragraph_format.space_before = Pt(54)
    set_run_font(version.add_run("版本日期  2026 年 10 月 12 日"), 10)
    document.add_page_break()


def add_contents(document: Document, headings: list[str]) -> None:
    heading = document.add_heading("目录与阅读建议", level=1)
    heading.paragraph_format.page_break_before = False
    paragraph = document.add_paragraph()
    add_inline(
        paragraph,
        "第 1 章建立全局结构，第 2 至 6 章逐项学习五类技能，第 7 章把全部技术连接成完整系统，"
        "第 8 至 10 章用于制定学习计划、校准简历措辞和自检。可使用 Word 导航窗格按标题跳转。",
    )
    for item in headings:
        entry = document.add_paragraph()
        entry.paragraph_format.left_indent = Inches(0.18)
        entry.paragraph_format.first_line_indent = Inches(-0.18)
        entry.paragraph_format.space_after = Pt(4)
        add_inline(entry, item)
    document.add_page_break()


def add_header(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run("AI Agent 工程专业技能系统讲解"), 8.5)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D9D9D9")
    borders = OxmlElement("w:pBdr")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def build_document() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    headings = [line[3:].strip() for line in lines if line.startswith("## ")]

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.35)

    configure_styles(document)
    add_cover(document)
    add_contents(document, headings)

    body_started = False
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()

        if not body_started:
            if lines[index].startswith("## "):
                body_started = True
            else:
                index += 1
                continue

        if stripped.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines: list[str] = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                table_lines.append(candidate)
                index += 1
            add_table(document, parse_table(table_lines))
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            markdown_level = len(heading_match.group(1))
            heading_text = heading_match.group(2)
            if markdown_level == 2:
                paragraph = document.add_heading(heading_text, level=1)
            else:
                paragraph = document.add_paragraph()
                add_inline(paragraph, heading_text)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(14 if markdown_level == 3 else 12)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.left_indent = Inches(0)
                paragraph.paragraph_format.right_indent = Inches(0)
                paragraph.paragraph_format.first_line_indent = Inches(0)
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.space_before = Pt(11)
                paragraph.paragraph_format.space_after = Pt(5)
            index += 1
            continue

        ordered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if ordered:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline(paragraph, f"{ordered.group(1)}.  {ordered.group(2)}")
            index += 1
            continue

        unordered = re.match(r"^-\s+(.+)$", stripped)
        if unordered:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, unordered.group(1))
            index += 1
            continue

        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip()[1:].strip())
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.right_indent = Inches(0.2)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(9)
            add_inline(paragraph, " ".join(quote_lines))
            for run in paragraph.runs:
                run.italic = True
            continue

        if not stripped:
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", "```", ">", "|", "- "))
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(paragraph, " ".join(paragraph_lines))

    for doc_section in document.sections:
        add_header(doc_section)
        footer_paragraph = doc_section.footer.paragraphs[0]
        add_page_number(footer_paragraph)

    document.core_properties.title = "AI Agent 工程专业技能系统讲解"
    document.core_properties.subject = "后端工程 Agent RAG 评测稳定性与 LLM 基础学习手册"
    document.core_properties.author = "WorkPilot"
    document.core_properties.keywords = (
        "Python, FastAPI, Agent, LangGraph, MCP, RAG, Evaluation, LLM"
    )
    document.core_properties.comments = "Generated from the maintained Markdown source."
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()

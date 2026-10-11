from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "newcomer-agent-learning-guide.md"
OUTPUT = ROOT / "docs" / "WorkPilot-新人学习与面试讲解手册.docx"

BODY_FONT = "Hiragino Sans GB"
EAST_ASIA_FONT = "Hiragino Sans GB"
MONO_FONT = "Arial Unicode MS"
BLACK = "000000"
DARK_BLUE = "17365D"
PALE_BLUE = "EEF4FA"
LIGHT_GRAY = "D9D9D9"
CODE_BG = "F3F4F6"


def set_run_font(run, size: float | None = None, *, bold: bool | None = None) -> None:
    run.font.name = BODY_FONT
    fonts = run._element.get_or_add_rPr().rFonts
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), EAST_ASIA_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(
    cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), LIGHT_GRAY)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char_begin, instr_text, fld_char_end))
    set_run_font(run, 9)


def add_inline(paragraph, text: str, *, size: float | None = None) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            fonts = run._element.get_or_add_rPr().rFonts
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{attribute}"), MONO_FONT)
            run.font.size = Pt((size or 10.5) - 0.5)
            run.font.color.rgb = RGBColor(0, 0, 0)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), CODE_BG)
            run._element.get_or_add_rPr().append(shading)
        else:
            link = re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", token)
            assert link is not None
            run = paragraph.add_run(f"{link.group(1)}（{link.group(2)}）")
            set_run_font(run, size)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size)


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    paragraph.paragraph_format.keep_together = True
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = MONO_FONT
        fonts = run._element.get_or_add_rPr().rFonts
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), MONO_FONT)
        run.font.size = Pt(8.8)
        run.font.color.rgb = RGBColor(32, 33, 36)
        if index < len(lines) - 1:
            run.add_break()


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        values = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(values)
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    if columns == 2:
        widths = [1.75, 5.15]
    elif columns == 3:
        widths = [1.35, 2.1, 3.45]
    elif columns == 4:
        widths = [0.8, 1.6, 2.5, 2.0]
    else:
        widths = [6.9 / columns] * columns

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[column_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
            if row_index == 0:
                set_cell_shading(cell, DARK_BLUE)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(value)
                run.font.name = BODY_FONT
                fonts = run._element.get_or_add_rPr().rFonts
                for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                    fonts.set(qn(f"w:{attribute}"), EAST_ASIA_FONT)
                run.font.size = Pt(9.2)
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if row_index % 2 == 0:
                    set_cell_shading(cell, PALE_BLUE)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                )
                add_inline(paragraph, value, size=9.2)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal_fonts = normal._element.get_or_add_rPr().rFonts
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal_fonts.set(qn(f"w:{attribute}"), EAST_ASIA_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = BODY_FONT
    title_fonts = title._element.get_or_add_rPr().rFonts
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        title_fonts.set(qn(f"w:{attribute}"), EAST_ASIA_FONT)
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(18)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    heading_sizes = {"Heading 1": 18, "Heading 2": 14, "Heading 3": 12}
    for name, size in heading_sizes.items():
        style = styles[name]
        style.font.name = BODY_FONT
        style_fonts = style._element.get_or_add_rPr().rFonts
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            style_fonts.set(qn(f"w:{attribute}"), EAST_ASIA_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        # Reset the built-in template's heading indents explicitly. LibreOffice can
        # otherwise carry a hanging indent across page boundaries and clip the
        # numeric prefix of a heading rendered at the top of a page.
        style.paragraph_format.left_indent = Inches(0)
        style.paragraph_format.right_indent = Inches(0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(16 if name == "Heading 1" else 11)
        style.paragraph_format.space_after = Pt(7 if name == "Heading 1" else 5)

    # Custom visual subheading styles avoid inheriting Word template numbering
    # or indentation while retaining consistent typography.
    manual_heading_styles = (
        ("Manual Heading 2", 14, True),
        ("Manual Heading 3", 12, True),
    )
    for name, size, keep_with_next in manual_heading_styles:
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.base_style = normal
        style.font.name = BODY_FONT
        style_fonts = style._element.get_or_add_rPr().rFonts
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            style_fonts.set(qn(f"w:{attribute}"), EAST_ASIA_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.left_indent = Inches(0)
        style.paragraph_format.right_indent = Inches(0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.keep_with_next = keep_with_next
        style.paragraph_format.keep_together = True
        style.paragraph_format.space_before = Pt(11)
        style.paragraph_format.space_after = Pt(5)

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = MONO_FONT
    code_fonts = code_style._element.get_or_add_rPr().rFonts
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        code_fonts.set(qn(f"w:{attribute}"), MONO_FONT)
    code_style.font.size = Pt(8.8)
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.right_indent = Inches(0.18)
    code_style.paragraph_format.space_before = Pt(5)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.05
    p_pr = code_style._element.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CODE_BG)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:color"), "A8BDD2")
    borders.append(left)
    p_pr.append(borders)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = BODY_FONT
        style_fonts = style._element.get_or_add_rPr().rFonts
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            style_fonts.set(qn(f"w:{attribute}"), EAST_ASIA_FONT)
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.2


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("WorkPilot 新人学习与面试讲解手册")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("面向 AI Agent 初学者的系统学习文档")
    set_run_font(run, 15)

    audience = document.add_paragraph()
    audience.alignment = WD_ALIGN_PARAGRAPH.CENTER
    audience.paragraph_format.space_after = Pt(8)
    run = audience.add_run("适用对象  刚接触 AI Agent 并需要接手 WorkPilot 的同学")
    set_run_font(run, 10.5)

    goal = document.add_paragraph()
    goal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    goal.paragraph_format.space_after = Pt(34)
    run = goal.add_run("学习目标  理解原理  读懂代码  独立演示  准确面试表达")
    set_run_font(run, 10.5)

    intro = document.add_paragraph()
    intro.paragraph_format.left_indent = Inches(0.55)
    intro.paragraph_format.right_indent = Inches(0.55)
    intro.paragraph_format.line_spacing = 1.35
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(
        intro,
        "本手册从生活化类比开始，逐步讲解 Agent Runtime、安全执行、持久工作流、"
        "多租户隔离、知识检索、可观测性和评测体系。完成学习后，读者应能亲自运行项目、"
        "解释关键设计、完成可验证的个人贡献，并在面试中准确区分项目能力与本人贡献。",
        size=11,
    )

    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(54)
    run = date.add_run("版本日期  2026 年 10 月 11 日")
    set_run_font(run, 10)
    document.add_page_break()


def add_contents(document: Document, headings: list[str]) -> None:
    heading = document.add_heading("目录与学习顺序", level=1)
    heading.paragraph_format.page_break_before = False
    paragraph = document.add_paragraph()
    add_inline(
        paragraph,
        "建议先完成第 1 至 3 章建立整体认识，再结合第 4 至 9 章阅读代码和动手实践，"
        "最后使用第 10 至 14 章准备面试并完成个人贡献。Word 导航窗格可按标题快速跳转。",
    )
    for number, item in enumerate(headings, start=1):
        entry = document.add_paragraph()
        entry.paragraph_format.left_indent = Inches(0.18)
        entry.paragraph_format.first_line_indent = Inches(-0.18)
        entry.paragraph_format.space_after = Pt(3)
        add_inline(entry, f"{number}.  {re.sub(r'^\d+\.\s*', '', item)}")
    document.add_page_break()


def build_document() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    headings = [line[3:].strip() for line in lines if line.startswith("## ")]

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    configure_styles(document)
    add_cover(document)
    add_contents(document, headings)

    body_started = False
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not body_started:
            if line.startswith("## "):
                body_started = True
            else:
                index += 1
                continue

        if stripped.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines: list[str] = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                table_lines.append(candidate)
                index += 1
            add_table(document, parse_table(table_lines))
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            markdown_level = len(heading_match.group(1))
            heading_text = heading_match.group(2)
            if markdown_level == 2:
                document.add_heading(heading_text, level=1)
            else:
                if heading_text.startswith("11."):
                    # Render the dense interview-question headings with direct
                    # formatting. Some LibreOffice builds shift custom heading
                    # styles past the left edge after a page transition.
                    paragraph = document.add_paragraph()
                    add_inline(paragraph, heading_text)
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.size = Pt(14)
                    paragraph.paragraph_format.keep_together = True
                    paragraph.paragraph_format.space_before = Pt(11)
                    paragraph.paragraph_format.space_after = Pt(5)
                else:
                    style = "Manual Heading 2" if markdown_level == 3 else "Manual Heading 3"
                    paragraph = document.add_paragraph(style=style)
                    add_inline(paragraph, heading_text)
                # Force the page-end questions onto fresh pages so each answer
                # starts alongside its question rather than on the next page.
                if heading_text.startswith(("11.7 ", "11.21 ")):
                    paragraph.paragraph_format.page_break_before = True
            index += 1
            continue

        ordered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if ordered:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, f"{ordered.group(1)}.  {ordered.group(2)}")
            index += 1
            continue

        unordered = re.match(r"^-\s+(.+)$", stripped)
        if unordered:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, unordered.group(1))
            index += 1
            continue

        if stripped.startswith(">"):
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip()[1:].strip())
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.right_indent = Inches(0.15)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            add_inline(paragraph, " ".join(quote_lines))
            for run in paragraph.runs:
                run.italic = True
            continue

        if not stripped:
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", "```", ">", "|", "- "))
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(paragraph, " ".join(paragraph_lines))

    for doc_section in document.sections:
        footer = doc_section.footer
        footer.is_linked_to_previous = True
        footer_paragraph = footer.paragraphs[0]
        add_page_number(footer_paragraph)

    document.core_properties.title = "WorkPilot 新人学习与面试讲解手册"
    document.core_properties.subject = "AI Agent 项目学习 技术栈与面试准备"
    document.core_properties.author = "WorkPilot"
    document.core_properties.keywords = "AI Agent, WorkPilot, Temporal, PostgreSQL, MCP, RAG"
    document.core_properties.comments = "Generated from the maintained Markdown learning guide."
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
